#!/usr/bin/env python3
"""
build_papers.py — Generates publication-grade IEEE-styled DOCX and PDF papers
for both English (PAPER_EN.md) and Arabic (PAPER_AR.md).
"""

import os
import re
import subprocess
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

WORKSPACE = "/home/mobta/university/malik"

# ─── XML & Styling Helpers ───────────────────────────────────────────────────

def set_rtl(paragraph, on=True):
    if on:
        paragraph._p.set(qn('xml:lang'), 'ar')
        rpr = paragraph._p.get_or_add_pPr()
        rpr.set(qn('w:bidi'), 'true')

def set_cell_rtl(cell):
    for p in cell.paragraphs:
        set_rtl(p, True)

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D5DD", sz="4"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_page_number_to_run(run):
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = "PAGE"
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    run._r.extend([fld1, instr, fld2, fld3])

def clean_text_and_math(text):
    """Normalize markdown math tokens and common symbols into clean Unicode."""
    text = text.replace(r"\le", "≤").replace(r"\ge", "≥").replace(r"\pm", "±").replace(r"\gg", "≫").replace(r"\ll", "≪")
    text = text.replace(r"\alpha", "α").replace(r"\in", "∈").replace(r"\text{ ms}", " ms").replace(r"\text{bytes}", " bytes")
    text = text.replace(r"t_{\text{policy}}", "t_policy").replace(r"t_{\\text{policy}}", "t_policy")
    text = text.replace(r"\text{", "").replace(r"}", "")
    text = text.replace(r"\\", "")
    text = re.sub(r"\$([^$]+)\$", r"\1", text)
    return text

def parse_inlines(text):
    """Parses text containing **bold**, *italic*, and `code` into runs."""
    pattern = re.compile(r'(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)')
    parts = pattern.split(text)
    runs = []
    for p in parts:
        if not p:
            continue
        if p.startswith('`') and p.endswith('`') and len(p) >= 2:
            runs.append((p[1:-1], False, False, True))
        elif p.startswith('**') and p.endswith('**') and len(p) >= 4:
            runs.append((p[2:-2], True, False, False))
        elif p.startswith('*') and p.endswith('*') and len(p) >= 2:
            runs.append((p[1:-1], False, True, False))
        else:
            runs.append((p, False, False, False))
    return runs

def add_styled_runs(para, text, rtl=False, default_size=10, default_font="Times New Roman", font_color=RGBColor(33, 37, 41)):
    cleaned = clean_text_and_math(text)
    runs_data = parse_inlines(cleaned)
    for t, bold, italic, code in runs_data:
        r = para.add_run(t)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(default_size)
        if code:
            r.font.name = "Consolas"
            r.font.size = Pt(default_size - 0.5)
            r.font.color.rgb = RGBColor(180, 40, 40)
        else:
            r.font.name = default_font if not rtl else "Noto Naskh Arabic"
            r.font.color.rgb = font_color

# ─── English Paper Builder ──────────────────────────────────────────────────

def build_paper_en():
    doc = Document()

    # Page Setup (US Letter, 0.75" margins)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # Header
    p_head = section.header.paragraphs[0]
    p_head.text = "Emirates International University | Faculty of Engineering & IT — Research Paper (IEEE Format)"
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_head.runs[0].font.name = "Times New Roman"
    p_head.runs[0].font.size = Pt(8.5)
    p_head.runs[0].font.italic = True
    p_head.runs[0].font.color.rgb = RGBColor(120, 120, 120)

    # Footer
    p_foot = section.footer.paragraphs[0]
    p_foot.text = "Evaluating Trust Boundaries in MCP Apps: Host-Mediated Capabilities  |  Page "
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.runs[0].font.name = "Times New Roman"
    p_foot.runs[0].font.size = Pt(8.5)
    p_foot.runs[0].font.color.rgb = RGBColor(120, 120, 120)
    add_page_number_to_run(p_foot.add_run())

    # Read Source Markdown
    md_path = os.path.join(WORKSPACE, "PAPER_EN.md")
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.splitlines()

    i = 0
    in_code_block = False
    code_lines = []
    code_lang = ""

    while i < len(lines):
        line = lines[i].rstrip()

        # Markdown Image Detection: ![alt](path)
        m_img = re.match(r"^!\[(.*?)\]\((.*?)\)$", line.strip())
        if m_img:
            alt_text = m_img.group(1)
            img_rel = m_img.group(2).strip()
            img_path = os.path.join(WORKSPACE, img_rel)
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.keep_with_next = True
                r_img = p_img.add_run()
                r_img.add_picture(img_path, width=Inches(6.8))
            i += 1
            continue

        # Code Block Delimiter
        if line.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = line[3:].strip()
                code_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                full_code = "\n".join(code_lines)
                # Skip rendering raw ASCII figure boxes if any remain
                if any(fig_kw in full_code for fig_kw in ["HOST POLICY ENFORCEMENT ENGINE", "LAB TESTBED TOPOLOGY", "HOST CLIENT"]):
                    i += 1
                    continue

                # Render code box
                box_table = doc.add_table(rows=1, cols=1)
                box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = box_table.rows[0].cells[0]
                set_cell_shading(cell, "F4F6F9")
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

                # XML left border highlight
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(r'''
                    <w:tcBorders {} >
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>
                        <w:left w:val="single" w:sz="18" w:space="0" w:color="1F497D"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>
                        <w:right w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>
                    </w:tcBorders>
                '''.format(nsdecls('w')))
                tcPr.append(tcBorders)

                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                cp.paragraph_format.line_spacing = Pt(10)
                r_code = cp.add_run(full_code)
                r_code.font.name = "Consolas"
                r_code.font.size = Pt(8)
                r_code.font.color.rgb = RGBColor(30, 41, 59)
                i += 1
                continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Skip horizontal dividers
        if line.strip() == "---":
            i += 1
            continue

        # Title
        if line.startswith("# "):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(title_text)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(18)
            r.font.color.rgb = RGBColor(10, 37, 64)
            i += 1
            continue

        # Authors / Metadata Block
        if line.startswith("**Authors:**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            add_styled_runs(p, line, default_size=10.5, font_color=RGBColor(33, 37, 41))
            i += 1
            continue

        if line.startswith("**Institution:**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            add_styled_runs(p, line, default_size=9.5, font_color=RGBColor(71, 85, 105))
            i += 1
            continue

        if line.startswith("**Date:**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            add_styled_runs(p, line, default_size=9.5, font_color=RGBColor(100, 116, 139))
            i += 1
            continue

        # Headings (IEEE Level 1: Roman numerals)
        if line.startswith("## "):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True

            if h_text.lower() == "abstract":
                r = p.add_run("ABSTRACT")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif h_text.lower() == "references":
                r = p.add_run("REFERENCES")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif h_text.lower().startswith("appendix"):
                r = p.add_run(h_text.upper())
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                m = re.match(r"^(\d+)\.\s*(.*)$", h_text)
                if m:
                    num_map = {1:"I", 2:"II", 3:"III", 4:"IV", 5:"V", 6:"VI", 7:"VII", 8:"VIII", 9:"IX"}
                    rom = num_map.get(int(m.group(1)), m.group(1))
                    r = p.add_run(f"{rom}.  {m.group(2).upper()}")
                else:
                    r = p.add_run(h_text.upper())
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11.5)
            r.font.color.rgb = RGBColor(15, 44, 89)
            i += 1
            continue

        # Headings (Level 2: A. Title)
        if line.startswith("### "):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True

            m = re.match(r"^\d+\.(\d+)\s*(.*)$", h_text)
            if m:
                sub_let = chr(64 + int(m.group(1))) if int(m.group(1)) <= 26 else m.group(1)
                r = p.add_run(f"{sub_let}.  {m.group(2)}")
            else:
                r = p.add_run(h_text)

            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.font.italic = True
            r.font.color.rgb = RGBColor(30, 58, 138)
            i += 1
            continue

        # Headings (Level 3: 1) Sub-title)
        if line.startswith("#### "):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1
            continue

        # Markdown Table Detection
        if line.startswith("|") and line.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 3:
                raw_headers = [c.strip() for c in table_lines[0].split("|")[1:-1]]
                raw_rows = []
                for tl in table_lines[2:]:
                    cols = [c.strip() for c in tl.split("|")[1:-1]]
                    if len(cols) == len(raw_headers):
                        raw_rows.append(cols)

                tbl = doc.add_table(rows=len(raw_rows) + 1, cols=len(raw_headers))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="CBD5E1", sz="6")

                # Format Header
                for col_idx, h_text in enumerate(raw_headers):
                    cell = tbl.rows[0].cells[col_idx]
                    set_cell_shading(cell, "1F497D")
                    set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
                    p_cell = cell.paragraphs[0]
                    p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cell.paragraph_format.space_after = Pt(0)
                    add_styled_runs(p_cell, h_text, default_size=9, font_color=RGBColor(255, 255, 255))
                    for r in p_cell.runs:
                        r.bold = True

                # Format Data Rows
                for row_idx, row_data in enumerate(raw_rows):
                    fill = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                    for col_idx, val in enumerate(row_data):
                        cell = tbl.rows[row_idx + 1].cells[col_idx]
                        if fill != "FFFFFF":
                            set_cell_shading(cell, fill)
                        set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
                        p_cell = cell.paragraphs[0]
                        p_cell.paragraph_format.space_after = Pt(0)
                        if re.match(r"^(\$?[0-9\.\+\-±\s]+(ms|%|leaks|delivered|executed|leaked|total|points)?\$?|N/A|Absent|Present)$", val, re.I):
                            p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            p_cell.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        add_styled_runs(p_cell, val, default_size=8.5)

                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.space_after = Pt(4)
            continue

        # Captions
        if line.startswith("*Figure ") or line.startswith("Fig. ") or line.startswith("*Fig"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(8)
            add_styled_runs(p, line, default_size=9, font_color=RGBColor(71, 85, 105))
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        if line.startswith("**Table ") or line.startswith("TABLE "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            add_styled_runs(p, line, default_size=9.5, font_color=RGBColor(15, 44, 89))
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        # Bullet List
        if line.startswith("* ") or line.startswith("- "):
            b_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_styled_runs(p, b_text, default_size=9.5)
            i += 1
            continue

        # Numbered List
        m_num = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m_num and not line.startswith("##"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            r_num = p.add_run(f"{m_num.group(1)}.  ")
            r_num.bold = True
            r_num.font.name = "Times New Roman"
            r_num.font.size = Pt(9.5)
            add_styled_runs(p, m_num.group(2), default_size=9.5)
            i += 1
            continue

        # Regular Paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            if line.startswith("The Model Context Protocol (MCP) standardizes how"):
                r_abs = p.add_run("Abstract—")
                r_abs.bold = True
                r_abs.italic = True
                r_abs.font.name = "Times New Roman"
                r_abs.font.size = Pt(10)
                add_styled_runs(p, line, default_size=9.5)
            elif line.startswith("**Keywords:**"):
                kw_text = line.replace("**Keywords:**", "").strip()
                r_kw = p.add_run("Index Terms—")
                r_kw.bold = True
                r_kw.italic = True
                r_kw.font.name = "Times New Roman"
                r_kw.font.size = Pt(9.5)
                add_styled_runs(p, kw_text, default_size=9.5)
            else:
                add_styled_runs(p, line, default_size=9.5)

        i += 1

    out_docx = os.path.join(WORKSPACE, "MCP_Apps_Trust_Boundaries_Paper_EN.docx")
    doc.save(out_docx)
    print(f"[OK] Generated: {out_docx}")
    return out_docx

# ─── Arabic Paper Builder ───────────────────────────────────────────────────

def build_paper_ar():
    doc = Document()

    # Page Setup (A4, 0.75" margins)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # Header
    p_head = section.header.paragraphs[0]
    set_rtl(p_head, True)
    p_head.text = "الجامعة الإماراتية الدولية | كلية الهندسة وتكنولوجيا المعلومات — ورقة بحثية محكمة"
    p_head.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_head.runs[0].font.name = "Noto Naskh Arabic"
    p_head.runs[0].font.size = Pt(8.5)
    p_head.runs[0].font.color.rgb = RGBColor(120, 120, 120)

    # Footer
    p_foot = section.footer.paragraphs[0]
    set_rtl(p_foot, True)
    p_foot.text = "تقييم حدود الثقة في تطبيقات MCP Apps  |  الصفحة "
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.runs[0].font.name = "Noto Naskh Arabic"
    p_foot.runs[0].font.size = Pt(8.5)
    p_foot.runs[0].font.color.rgb = RGBColor(120, 120, 120)
    add_page_number_to_run(p_foot.add_run())

    # Read Source Markdown
    md_path = os.path.join(WORKSPACE, "PAPER_AR.md")
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    lines = content.splitlines()

    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i].rstrip()

        # Markdown Image Detection: ![alt](path)
        m_img = re.match(r"^!\[(.*?)\]\((.*?)\)$", line.strip())
        if m_img:
            alt_text = m_img.group(1)
            img_rel = m_img.group(2).strip()
            img_path = os.path.join(WORKSPACE, img_rel)
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                p_img.paragraph_format.keep_with_next = True
                r_img = p_img.add_run()
                r_img.add_picture(img_path, width=Inches(6.8))
            i += 1
            continue

        # Code Block Delimiter
        if line.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = line[3:].strip()
                code_lines = []
                i += 1
                continue
            else:
                in_code_block = False
                full_code = "\n".join(code_lines)
                if any(fig_kw in full_code for fig_kw in ["HOST POLICY ENFORCEMENT ENGINE", "LAB TESTBED TOPOLOGY", "HOST CLIENT", "محرك سياسة المضيف المقترح", "تطبيق المضيف", "طوبولوجيا"]):
                    i += 1
                    continue

                box_table = doc.add_table(rows=1, cols=1)
                box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = box_table.rows[0].cells[0]
                set_cell_shading(cell, "F4F6F9")
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(r'''
                    <w:tcBorders {} >
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>
                        <w:left w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/>
                        <w:right w:val="single" w:sz="18" w:space="0" w:color="1F497D"/>
                    </w:tcBorders>
                '''.format(nsdecls('w')))
                tcPr.append(tcBorders)

                cp = cell.paragraphs[0]
                set_rtl(cp, False)  # diagrams are LTR monospaced
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                cp.paragraph_format.line_spacing = Pt(10)
                r_code = cp.add_run(full_code)
                r_code.font.name = "Consolas"
                r_code.font.size = Pt(8)
                r_code.font.color.rgb = RGBColor(30, 41, 59)
                i += 1
                continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Skip horizontal dividers
        if line.strip() == "---":
            i += 1
            continue

        # Title
        if line.startswith("# "):
            title_text = line[2:].strip()
            p = doc.add_paragraph()
            set_rtl(p, True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(title_text)
            r.bold = True
            r.font.name = "Noto Naskh Arabic"
            r.font.size = Pt(18)
            r.font.color.rgb = RGBColor(10, 37, 64)
            i += 1
            continue

        if line.startswith("## Evaluating Trust Boundaries"):
            sub_title = line[2:].strip()
            p = doc.add_paragraph()
            set_rtl(p, False)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(sub_title)
            r.italic = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Authors / Metadata Block
        if line.startswith("**الباحثون:**") or line.startswith("**الجهة:**") or line.startswith("**التاريخ:**"):
            p = doc.add_paragraph()
            set_rtl(p, True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3)
            add_styled_runs(p, line, rtl=True, default_size=10.5, font_color=RGBColor(33, 37, 41))
            i += 1
            continue

        # Headings Level 1
        if line.startswith("## "):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            set_rtl(p, True)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_styled_runs(p, h_text, rtl=True, default_size=12, font_color=RGBColor(15, 44, 89))
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        # Headings Level 2
        if line.startswith("### "):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            set_rtl(p, True)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_styled_runs(p, h_text, rtl=True, default_size=11, font_color=RGBColor(30, 58, 138))
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        # Headings Level 3
        if line.startswith("#### "):
            h_text = line[5:].strip()
            p = doc.add_paragraph()
            set_rtl(p, True)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_styled_runs(p, h_text, rtl=True, default_size=10.5, font_color=RGBColor(51, 65, 85))
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        # Markdown Table Detection
        if line.startswith("|") and line.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 3:
                raw_headers = [c.strip() for c in table_lines[0].split("|")[1:-1]]
                raw_rows = []
                for tl in table_lines[2:]:
                    cols = [c.strip() for c in tl.split("|")[1:-1]]
                    if len(cols) == len(raw_headers):
                        raw_rows.append(cols)

                tbl = doc.add_table(rows=len(raw_rows) + 1, cols=len(raw_headers))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl, color="CBD5E1", sz="6")

                # Format Header
                for col_idx, h_text in enumerate(raw_headers):
                    cell = tbl.rows[0].cells[col_idx]
                    set_cell_shading(cell, "1F497D")
                    set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
                    p_cell = cell.paragraphs[0]
                    set_rtl(p_cell, True)
                    p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cell.paragraph_format.space_after = Pt(0)
                    add_styled_runs(p_cell, h_text, rtl=True, default_size=9.5, font_color=RGBColor(255, 255, 255))
                    for r in p_cell.runs:
                        r.bold = True

                # Format Data Rows
                for row_idx, row_data in enumerate(raw_rows):
                    fill = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
                    for col_idx, val in enumerate(row_data):
                        cell = tbl.rows[row_idx + 1].cells[col_idx]
                        if fill != "FFFFFF":
                            set_cell_shading(cell, fill)
                        set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
                        p_cell = cell.paragraphs[0]
                        set_rtl(p_cell, True)
                        p_cell.paragraph_format.space_after = Pt(0)
                        if re.match(r"^(\$?[0-9\.\+\-±\s]+(مللي ثانية|%|تسريبات|نُفذت|وصلت|خروقات)?\$?|لا ينطبق|غائبة|ظاهرة)$", val):
                            p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            p_cell.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        add_styled_runs(p_cell, val, rtl=True, default_size=9)

                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.space_after = Pt(4)
            continue

        # Captions
        if line.startswith("*الشكل ") or line.startswith("الشكل "):
            p = doc.add_paragraph()
            set_rtl(p, True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(8)
            add_styled_runs(p, line, rtl=True, default_size=9.5, font_color=RGBColor(71, 85, 105))
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        if line.startswith("**الجدول ") or line.startswith("الجدول "):
            p = doc.add_paragraph()
            set_rtl(p, True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            add_styled_runs(p, line, rtl=True, default_size=10, font_color=RGBColor(15, 44, 89))
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        # Bullet List
        if line.startswith("* ") or line.startswith("- "):
            b_text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            set_rtl(p, True)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_styled_runs(p, b_text, rtl=True, default_size=10)
            i += 1
            continue

        # Numbered List
        m_num = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m_num and not line.startswith("##"):
            p = doc.add_paragraph()
            set_rtl(p, True)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_num = p.add_run(f"{m_num.group(1)}.  ")
            r_num.bold = True
            r_num.font.name = "Noto Naskh Arabic"
            r_num.font.size = Pt(10)
            add_styled_runs(p, m_num.group(2), rtl=True, default_size=10)
            i += 1
            continue

        # Regular Paragraph
        if line.strip():
            p = doc.add_paragraph()
            set_rtl(p, True)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            if line.startswith("يُعد بروتوكول سياق النموذج (Model Context Protocol"):
                r_abs = p.add_run("الملخص الأكاديمي — ")
                r_abs.bold = True
                r_abs.font.name = "Noto Naskh Arabic"
                r_abs.font.size = Pt(10.5)
                add_styled_runs(p, line, rtl=True, default_size=10)
            elif line.startswith("**الكلمات المفتاحية:**"):
                kw_text = line.replace("**الكلمات المفتاحية:**", "").strip()
                r_kw = p.add_run("الكلمات المفتاحية — ")
                r_kw.bold = True
                r_kw.font.name = "Noto Naskh Arabic"
                r_kw.font.size = Pt(10)
                add_styled_runs(p, kw_text, rtl=True, default_size=10)
            else:
                add_styled_runs(p, line, rtl=True, default_size=10)

        i += 1

    out_docx = os.path.join(WORKSPACE, "MCP_Apps_Trust_Boundaries_Paper_AR.docx")
    doc.save(out_docx)
    print(f"[OK] Generated: {out_docx}")
    return out_docx

# ─── PDF Conversion ──────────────────────────────────────────────────────────

def convert_to_pdf(docx_path):
    print(f"Converting {os.path.basename(docx_path)} to PDF...")
    cmd = ["libreoffice", "--headless", "--convert-to", "pdf", docx_path]
    res = subprocess.run(cmd, cwd=WORKSPACE, capture_output=True, text=True)
    if res.returncode == 0:
        pdf_name = os.path.splitext(docx_path)[0] + ".pdf"
        print(f"[OK] Generated PDF: {pdf_name}")
        return pdf_name
    else:
        print(f"[ERR] Failed to convert {docx_path} to PDF: {res.stderr}")
        return None

# ─── Main Execution ──────────────────────────────────────────────────────────

if __name__ == "__main__":
    en_docx = build_paper_en()
    ar_docx = build_paper_ar()
    convert_to_pdf(en_docx)
    convert_to_pdf(ar_docx)
    print("\nAll academic research paper artifacts generated successfully!")
